import codecs
import csv
import datetime
import decimal
import io
import json
import math
import os
import pickle
import re
import time
import duckdb

class TypedCopy:

    def __init__(self, path, wanted=True):
        self.path, self.file, self.finished = (path, None, None)
        if wanted:
            try:
                self.file = open(path, 'wb')
            except OSError:
                self.drop()
        else:
            self.drop()

    def add(self, batch):
        if self.file is not None:
            try:
                pickle.dump(batch, self.file)
            except Exception:
                self.drop()

    def drop(self):
        file, self.file, self.finished = (self.file, None, None)
        try:
            if file is not None:
                file.close()
        except Exception:
            pass
        try:
            if os.path.exists(self.path):
                os.remove(self.path)
        except OSError:
            pass

    def finish(self):
        file, self.file = (self.file, None)
        if file is not None:
            try:
                file.close()
                self.finished = self.path
            except Exception:
                self.drop()
        return self.finished

class DuckDBCSVTools:
    PREVIEW_ROWS = 1000
    XLSX_ROWS = 1048575
    XLSX_COLUMNS = 16384
    XLSX_TEXT = 32767
    XLSX_DIGITS = 15
    XLSX_LARGEST = 9.99999999999999e+307
    XLSX_SMALLEST = 2.2250738585072014e-308
    CLIPBOARD_CHARS = 10000000
    WORKBOOK_PATH = '/tmp/duck_export.xlsx'
    PARQUET_PATH = '/tmp/duck_export.parquet'

    def __init__(self):
        self.conn = duckdb.connect(':memory:')
        self.tables = {}
        self.sources = {}
        self.counter = 0
        self.result_id = 0
        self.snapshot = None

    def quote_identifier(self, value):
        if not isinstance(value, str) or not value:
            raise ValueError('SQL identifier must be a non-empty string')
        return '"' + value.replace('"', '""') + '"'

    def current_namespace(self):
        return (self.conn.execute('SELECT current_database()').fetchone()[0], 'main')

    def relation_key(self, name):
        return (*self.current_namespace(), name)

    def quote_relation(self, name):
        return '.'.join((self.quote_identifier(part) for part in self.relation_key(name)))

    def quote_literal(self, value):
        if not isinstance(value, str):
            raise TypeError('SQL literal must be a string')
        return "'" + value.replace("'", "''") + "'"

    def get_version(self):
        return {'duckdb': duckdb.__version__, 'preview_rows': self.PREVIEW_ROWS}

    def identifier_key(self, name):
        return re.sub('[A-Z]', lambda match: match.group().lower(), name)

    def new_name(self, filename, sheet=None, requested_name=None, reserved_names=None):
        stem = os.path.splitext(filename)[0]
        if sheet:
            stem += '_' + sheet
        stem = re.sub('[^a-zA-Z0-9_]', '_', stem).strip('_').lower() or 'table'
        if stem[0].isdigit():
            stem = 'table_' + stem
        names = {self.identifier_key(name) for name in self.relation_catalog()}
        names.update((self.identifier_key(row[0]) for row in self.conn.execute('SHOW TABLES').fetchall()))
        if requested_name is not None:
            self.quote_identifier(requested_name)
            if len(requested_name) > 1000 or self.identifier_key(requested_name) in names:
                raise ValueError('Saved table name is already in use or too long')
            self.counter += 1
            return (requested_name, self.counter)
        names.update((self.identifier_key(name) for name in reserved_names or []))
        name, suffix = (stem, 2)
        while self.identifier_key(name) in names:
            name = f'{stem}_{suffix}'
            suffix += 1
        self.counter += 1
        return (name, self.counter)

    def relation_catalog(self):
        rows = self.conn.execute("\n            SELECT view_name, 'VIEW', view_oid FROM duckdb_views()\n            WHERE database_name = current_database() AND schema_name = 'main' AND NOT internal\n            UNION ALL\n            SELECT table_name, 'BASE TABLE', table_oid FROM duckdb_tables()\n            WHERE database_name = current_database() AND schema_name = 'main' AND NOT internal\n        ").fetchall()
        return {name: (kind, oid) for name, kind, oid in rows}

    def table_metadata(self, name, info, count_rows=True):
        quoted = self.quote_relation(name)
        if count_rows:
            schema = self.conn.execute(f'DESCRIBE {quoted}').fetchall()
        else:
            schema = self.conn.execute("SELECT column_name, data_type FROM duckdb_columns() WHERE database_name = current_database() AND schema_name = 'main' AND table_name = ? ORDER BY column_index", [name]).fetchall()
        if info['file_type'] != 'sql':
            info = dict(info, relation_oid=self.relation_catalog()[name][1])
        return dict(info, database=self.current_namespace()[0], schema='main', columns=[row[0] for row in schema], types=[row[1] for row in schema], rows=self.conn.execute(f'SELECT count(*) FROM {quoted}').fetchone()[0] if count_rows else None)

    def csv_select(self, path, options):
        header = 'true' if options.get('has_headers', True) else 'false'
        text = 'true' if options.get('all_varchar', False) else 'false'
        delimiter = options.get('delimiter', '')
        if delimiter and (not isinstance(delimiter, str) or len(delimiter) != 1):
            raise ValueError('Delimiter must be one character')
        delim_sql = f', delim={self.quote_literal(delimiter)}' if delimiter else ''
        return f'read_csv({self.quote_literal(path)}, header={header}, auto_detect=true, all_varchar={text}{delim_sql})'

    def check_utf8(self, path):
        decoder = codecs.getincrementaldecoder('utf-8')()
        try:
            with open(path, 'rb') as file:
                while (chunk := file.read(1 << 20)):
                    decoder.decode(chunk)
            decoder.decode(b'', final=True)
        except UnicodeDecodeError:
            raise ValueError('This CSV is not UTF-8 text. Save it as CSV UTF-8 and try again.') from None

    def load_csv_file(self, upload, filename, options=None, requested_name=None, reserved_names=None, origin='file'):
        path = None
        created = False
        try:
            if origin not in ('file', 'sample'):
                raise ValueError('Invalid source origin')
            options = options or {}
            name, source_id = self.new_name(filename, requested_name=requested_name, reserved_names=reserved_names)
            path = f'/tmp/duck_source_{source_id}.csv'
            os.replace(upload, path)
            self.check_utf8(path)
            self.conn.execute(f'CREATE VIEW {self.quote_relation(name)} AS SELECT * FROM {self.csv_select(path, options)}')
            created = True
            info = self.table_metadata(name, {'filename': filename, 'file_type': 'csv', 'origin': origin, 'source_id': source_id, 'options': options, 'has_headers': options.get('has_headers', True)})
            self.tables[self.relation_key(name)] = info
            self.sources[source_id] = {'path': path}
            return dict(info, success=True, table_name=name)
        except Exception as error:
            if path:
                if created:
                    self.conn.execute(f'DROP VIEW IF EXISTS {self.quote_relation(name)}')
                if os.path.exists(path):
                    os.remove(path)
            return {'success': False, 'error': str(error)}

    def load_excel_file(self, upload, filename, requested_name=None, requested_sheet=None, reserved_names=None):
        try:
            import pandas as pd
            with pd.ExcelFile(upload, engine='openpyxl') as book:
                sheets = book.sheet_names
                if not sheets:
                    raise ValueError('Workbook has no worksheets')
                sheet = requested_sheet or sheets[0]
                if sheet not in sheets:
                    raise ValueError(f'Saved worksheet {sheet!r} is not in this workbook')
                frame = pd.read_excel(book, sheet_name=sheet)
            name, source_id = self.new_name(filename, requested_name=requested_name, reserved_names=reserved_names)
            self.create_excel_table(name, frame)
            info = self.table_metadata(name, {'filename': filename, 'file_type': 'excel', 'origin': 'file', 'source_id': source_id, 'available_sheets': sheets, 'current_sheet': sheet})
            path = f'/tmp/duck_source_{source_id}.xlsx'
            os.replace(upload, path)
            self.tables[self.relation_key(name)] = info
            self.sources[source_id] = {'path': path}
            return dict(info, success=True, table_name=name)
        except Exception as error:
            return {'success': False, 'error': str(error)}

    def create_excel_table(self, name, frame):
        registration = '_duck_import_frame'
        self.conn.register(registration, frame)
        try:
            self.conn.execute(f'CREATE TABLE {self.quote_relation(name)} AS SELECT * FROM {registration}')
        finally:
            self.conn.unregister(registration)

    def load_additional_sheet(self, base_name, sheet, requested_name=None, reserved_names=None):
        try:
            import pandas as pd
            self.get_table_info()
            base = self.tables[self.relation_key(base_name)]
            source_id = base['source_id']
            if sheet not in base.get('available_sheets', []):
                raise ValueError('Worksheet not found')
            if any((info['source_id'] == source_id and info.get('current_sheet') == sheet for info in self.tables.values())):
                raise ValueError('Worksheet is already loaded')
            frame = pd.read_excel(self.sources[source_id]['path'], sheet_name=sheet, engine='openpyxl')
            name, _ = self.new_name(base['filename'], sheet, requested_name, reserved_names)
            self.create_excel_table(name, frame)
            info = self.table_metadata(name, dict(base, current_sheet=sheet))
            self.tables[self.relation_key(name)] = info
            return dict(info, success=True, table_name=name)
        except Exception as error:
            return {'success': False, 'error': str(error)}

    def reload_table_with_headers(self, name, has_headers):
        try:
            self.get_table_info()
            info = self.tables[self.relation_key(name)]
            if info['file_type'] != 'csv':
                raise ValueError('Header settings apply to CSV files')
            options = dict(info['options'], has_headers=bool(has_headers))
            select = self.csv_select(self.sources[info['source_id']]['path'], options)
            self.conn.execute(f'SELECT count(*) FROM {select}').fetchone()
            self.conn.execute('BEGIN TRANSACTION')
            try:
                self.conn.execute(f'CREATE OR REPLACE VIEW {self.quote_relation(name)} AS SELECT * FROM {select}')
                refreshed = self.table_metadata(name, dict(info, options=options, has_headers=bool(has_headers)))
                self.conn.execute('COMMIT')
            except Exception:
                self.conn.execute('ROLLBACK')
                raise
            self.tables[self.relation_key(name)] = refreshed
            return dict(refreshed, success=True, table_name=name)
        except Exception as error:
            return {'success': False, 'error': str(error)}

    def replace_file(self, upload, name, filename):
        path = None
        try:
            self.get_table_info()
            info = self.tables.get(self.relation_key(name))
            if not info or info.get('origin') != 'file':
                raise ValueError('Only a loaded file you added can be replaced')
            self.counter += 1
            source_id = self.counter
            registration = None
            if info['file_type'] == 'csv':
                if not filename.lower().endswith('.csv'):
                    raise ValueError('Replace a CSV source with a CSV file')
                path = f'/tmp/duck_source_{source_id}.csv'
                os.replace(upload, path)
                self.check_utf8(path)
                select = self.csv_select(path, info['options'])
                self.conn.execute(f'SELECT count(*) FROM {select}').fetchone()
                statement = f'CREATE OR REPLACE VIEW {self.quote_relation(name)} AS SELECT * FROM {select}'
                source, details = ({'path': path}, {})
            else:
                import pandas as pd
                if not filename.lower().endswith('.xlsx'):
                    raise ValueError('Replace an Excel source with an Excel file')
                sheet = info['current_sheet']
                path = f'/tmp/duck_source_{source_id}.xlsx'
                os.replace(upload, path)
                with pd.ExcelFile(path, engine='openpyxl') as book:
                    sheets = book.sheet_names
                    if sheet not in sheets:
                        raise ValueError(f'Saved worksheet {sheet!r} is not in this workbook')
                    frame = pd.read_excel(book, sheet_name=sheet)
                registration = '_duck_import_frame'
                self.conn.register(registration, frame)
                statement = f'CREATE OR REPLACE TABLE {self.quote_relation(name)} AS SELECT * FROM {registration}'
                source, details = ({'path': path}, {'available_sheets': sheets})
            try:
                self.conn.execute('BEGIN TRANSACTION')
                try:
                    self.conn.execute(statement)
                    refreshed = self.table_metadata(name, dict(info, filename=filename, source_id=source_id, **details))
                    self.conn.execute('COMMIT')
                except Exception:
                    try:
                        self.conn.execute('ROLLBACK')
                    except Exception:
                        pass
                    raise
            finally:
                if registration:
                    self.conn.unregister(registration)
            self.tables[self.relation_key(name)] = refreshed
            self.sources[source_id] = source
            old_id = info['source_id']
            if not any((item['source_id'] == old_id for item in self.tables.values())):
                old = self.sources.pop(old_id, {})
                try:
                    if 'path' in old and os.path.exists(old['path']):
                        os.remove(old['path'])
                except OSError:
                    pass
            return dict(refreshed, success=True, table_name=name)
        except Exception as error:
            if path and path not in (item.get('path') for item in self.sources.values()) and os.path.exists(path):
                os.remove(path)
            return {'success': False, 'error': str(error)}

    def drop_table(self, name):
        try:
            self.get_table_info()
            info = self.tables[self.relation_key(name)]
            kind = 'VIEW' if info['file_type'] == 'csv' else 'TABLE'
            self.conn.execute(f'DROP {kind} {self.quote_relation(name)}')
            del self.tables[self.relation_key(name)]
            source_id = info['source_id']
            if not any((item['source_id'] == source_id for item in self.tables.values())):
                os.remove(self.sources.pop(source_id)['path'])
            return {'success': True}
        except Exception as error:
            return {'success': False, 'error': str(error)}

    def get_table_info(self):
        namespace = self.current_namespace()
        catalog = self.relation_catalog()
        catalog_kinds = {name: identity[0] for name, identity in catalog.items()}
        for key, info in list(self.tables.items()):
            if key[:2] != namespace:
                continue
            expected_kind = 'VIEW' if info['file_type'] == 'csv' else 'BASE TABLE'
            if catalog.get(key[2]) != (expected_kind, info['relation_oid']):
                del self.tables[key]
        tables = {}
        for name, kind in catalog_kinds.items():
            info = self.tables.get((*namespace, name), {'filename': 'Created with SQL', 'file_type': 'sql'})
            tables[name] = info if info['file_type'] == 'csv' else self.table_metadata(name, info, count_rows=kind != 'VIEW')
        return {'tables': tables, 'count': len(tables), 'database': namespace[0]}

    def json_cell(self, value):
        if isinstance(value, int) and (not isinstance(value, bool)) and (abs(value) > 9007199254740991):
            return str(value)
        if isinstance(value, (decimal.Decimal, datetime.date, datetime.time)):
            return str(value)
        if isinstance(value, float) and (not math.isfinite(value)):
            return str(value)
        if isinstance(value, bytes):
            return value.hex()
        if isinstance(value, (tuple, list)):
            return [self.json_cell(item) for item in value]
        if isinstance(value, dict):
            return {str(key): self.json_cell(item) for key, item in value.items()}
        if value is None or isinstance(value, (str, bool, int, float)):
            return value
        return str(value)

    def neutralize_spreadsheet_formula(self, value):
        if value is None:
            return ''
        text = str(value)
        if text.startswith(('=', '+', '-', '@', '\t', '\r', '\n', '＝', '＋', '－', '＠')):
            return "'" + text
        return value

    def preview_query(self, sql, limit=100):
        started = time.perf_counter()
        try:
            if type(limit) is not int or not 1 <= limit <= self.PREVIEW_ROWS:
                raise ValueError('Preview limit must be an integer from 1 to 1000.')
            self.require_single_select(sql, 'Preview accepts one SELECT query. Use the SQL workbench for other statements.')
            relation = self.conn.sql(sql).limit(limit + 1)
            columns, types = (relation.columns, [str(value) for value in relation.types])
            rows = relation.fetchall()
            data = [[self.json_cell(value) for value in row] for row in rows[:limit]]
            return {'success': True, 'columns': columns, 'types': types, 'data': data, 'rows': len(data), 'truncated': len(rows) > limit, 'query': sql, 'elapsed_ms': round((time.perf_counter() - started) * 1000)}
        except Exception as error:
            return {'success': False, 'error': str(error)}

    def require_single_select(self, sql, message):
        statements = self.conn.extract_statements(sql)
        if len(statements) != 1 or statements[0].type != duckdb.StatementType.SELECT:
            raise ValueError(message)

    def execute_recipe(self, sql):
        return self._run(sql, single_select=True)

    def check_stages(self, stages):
        if not isinstance(stages, list) or not stages or (not all((isinstance(item, list) and len(item) == 2 and isinstance(item[1], str) and (isinstance(item[0], str) or (index == 0 and item[0] is None)) for index, item in enumerate(stages)))):
            raise ValueError('A staged recipe needs one name and query per step.')
        names = {self.identifier_key(name) for name in self.relation_catalog()}
        for name, _ in stages[:-1]:
            if name is not None and (not re.fullmatch('__+duck_stage_[0-9]+', name) or self.identifier_key(name) in names):
                raise ValueError('Invalid recipe stage name.')

    def execute_recipe_stages(self, stages):
        created, counts, began, committed = ([], [], False, False)
        try:
            self.conn.execute('BEGIN TRANSACTION')
            began = True
            self.check_stages(stages)
            for index, (name, sql) in enumerate(stages[:-1]):
                try:
                    self.require_single_select(sql, 'A recipe runs as one SELECT query per step. Use the SQL workbench for other statements.')
                    if name is None:
                        counts.append(self.conn.execute(f'SELECT count(*) FROM ({sql}) AS input').fetchone()[0])
                        continue
                    quoted = '"' + name + '"'
                    self.conn.execute(f'CREATE OR REPLACE TEMP TABLE {quoted} AS {sql}')
                    created.append(quoted)
                    counts.append(self.conn.execute(f'SELECT count(*) FROM temp.main.{quoted}').fetchone()[0])
                    while len(created) > 1:
                        self.conn.execute(f'DROP TABLE IF EXISTS temp.main.{created[0]}')
                        created.pop(0)
                except Exception as error:
                    return {'success': False, 'error': str(error), 'stage': index}

            def finish():
                nonlocal committed, finishing
                finishing = True
                for quoted in created:
                    self.conn.execute(f'DROP TABLE IF EXISTS temp.main.{quoted}')
                created.clear()
                self.conn.execute('COMMIT')
                committed = True
            finishing = False
            result = self._run(stages[-1][1], single_select=True, finish=finish)
            if not result['success']:
                if not finishing:
                    result['stage'] = len(stages) - 1
                return result
            result['counts'] = counts + [result['rows']]
            return result
        except Exception as error:
            return {'success': False, 'error': str(error)}
        finally:
            if began and (not committed):
                try:
                    self.conn.execute('ROLLBACK')
                except Exception:
                    pass

    def execute_query(self, sql):
        return self._run(sql)

    def _run(self, sql, single_select=False, finish=None):
        started = time.perf_counter()
        path = f'/tmp/duck_result_{self.result_id + 1}.csv'
        typed = f'/tmp/duck_result_{self.result_id + 1}.rows'
        copy = None
        try:
            if single_select:
                self.require_single_select(sql, 'A recipe runs as one SELECT query. Use the SQL workbench for other statements.')
            cursor = self.conn.execute(sql)
            columns = [item[0] for item in cursor.description or []]
            types = [str(item[1]) for item in cursor.description or []]
            preview, count = ([], 0)
            copy = TypedCopy(typed, wanted=len(columns) <= self.XLSX_COLUMNS)
            with open(path, 'w', encoding='utf-8', newline='') as output:
                writer = csv.writer(output, lineterminator='\r\n')
                writer.writerow([self.neutralize_spreadsheet_formula(name) for name in columns])
                while True:
                    batch = cursor.fetchmany(512)
                    if not batch:
                        break
                    if count + len(batch) > self.XLSX_ROWS:
                        copy.drop()
                    copy.add(batch)
                    for row in batch:
                        if count < self.PREVIEW_ROWS:
                            preview.append([self.json_cell(value) for value in row])
                        writer.writerow([self.neutralize_spreadsheet_formula(value) for value in row])
                        count += 1
            if finish is not None:
                finish()
            previous = self.snapshot
            self.result_id += 1
            self.snapshot = {'path': path, 'typed': copy.finish(), 'columns': columns, 'rows': count, 'id': self.result_id}
            self.discard_snapshot(previous)
            return {'success': True, 'columns': columns, 'types': types, 'data': preview, 'rows': count, 'truncated': count > self.PREVIEW_ROWS, 'query': sql, 'result_id': self.result_id, 'elapsed_ms': round((time.perf_counter() - started) * 1000)}
        except Exception as error:
            if copy is not None:
                copy.drop()
            self.discard_snapshot({'path': path})
            return {'success': False, 'error': str(error)}

    def discard_snapshot(self, snapshot):
        snapshot = snapshot or {}
        for name in (snapshot.get('path'), snapshot.get('typed')):
            try:
                if name and os.path.exists(name):
                    os.remove(name)
            except OSError:
                pass

    def excel_number(self, number):
        return number == 0 or (math.isfinite(number) and self.XLSX_SMALLEST <= abs(number) <= self.XLSX_LARGEST)

    def excel_cell(self, sheet, value, string_cell, illegal_characters, error_codes):
        if value is None or isinstance(value, bool):
            return value
        if isinstance(value, int):
            return value if abs(value) < 10 ** self.XLSX_DIGITS else str(value)
        if isinstance(value, float):
            return value if self.excel_number(value) else str(value)
        if isinstance(value, decimal.Decimal):
            return float(value) if value.is_finite() and len(value.as_tuple().digits) <= self.XLSX_DIGITS else str(value)
        if isinstance(value, datetime.datetime):
            return value if value.tzinfo is None and value.year >= 1900 and (value.microsecond % 1000 == 0) else value.isoformat(sep=' ')
        if isinstance(value, datetime.date):
            return value if value.year >= 1900 else value.isoformat()
        if isinstance(value, datetime.time):
            return value if value.tzinfo is None and value.microsecond % 1000 == 0 else value.isoformat()
        text = value.hex() if isinstance(value, bytes) else value if isinstance(value, str) else str(value)
        text = illegal_characters.sub('', text)
        if len(text) > self.XLSX_TEXT:
            raise ValueError(f'A cell holds more than {self.XLSX_TEXT:,} characters, which Excel cannot store. Use Download CSV instead.')
        if text.startswith('=') or text in error_codes:
            cell = string_cell(sheet, value=text)
            cell.data_type = 's'
            return cell
        return text

    def check_export_to_xlsx(self, result_id):
        snapshot = self.snapshot
        if not snapshot or snapshot['id'] != result_id:
            return {'success': False, 'error': 'Result is no longer available. Run the query again.'}
        if snapshot['rows'] > self.XLSX_ROWS or len(snapshot['columns']) > self.XLSX_COLUMNS:
            return {'success': False, 'error': f"This result has {snapshot['rows']:,} rows and {len(snapshot['columns']):,} columns. One Excel sheet holds {self.XLSX_ROWS:,} rows below the header and {self.XLSX_COLUMNS:,} columns. Use Download CSV instead."}
        if not snapshot['typed']:
            return {'success': False, 'error': 'This result could not be prepared for Excel. Use Download CSV instead.'}
        return {'success': True}

    def export_to_xlsx(self, result_id):
        check = self.check_export_to_xlsx(result_id)
        if not check['success']:
            return check
        snapshot = self.snapshot
        book = None
        try:
            from openpyxl import Workbook
            from openpyxl.cell import WriteOnlyCell
            from openpyxl.cell.cell import ERROR_CODES, ILLEGAL_CHARACTERS_RE
            book = Workbook(write_only=True)
            sheet = book.create_sheet('Result')
            cell = lambda value: self.excel_cell(sheet, value, WriteOnlyCell, ILLEGAL_CHARACTERS_RE, ERROR_CODES)
            sheet.append([cell(str(name)) for name in snapshot['columns']])
            with open(snapshot['typed'], 'rb') as rows:
                while True:
                    try:
                        batch = pickle.load(rows)
                    except EOFError:
                        break
                    for row in batch:
                        sheet.append([cell(value) for value in row])
            book.save(self.WORKBOOK_PATH)
            return {'success': True, 'send_file': self.send_file(self.WORKBOOK_PATH, remove=True), 'rows': snapshot['rows']}
        except Exception as error:
            self.discard_workbook(book)
            return {'success': False, 'error': str(error) or f'{type(error).__name__} while building the workbook. Use Download CSV instead.'}

    def send_file(self, path, remove):
        return {'path': path, 'remove': remove}

    def discard_workbook(self, book):
        try:
            if os.path.exists(self.WORKBOOK_PATH):
                os.remove(self.WORKBOOK_PATH)
        except OSError:
            pass
        for sheet in book.worksheets if book is not None else ():
            writer = getattr(sheet, '_writer', None)
            if writer is not None:
                for step in (sheet.close, writer.cleanup):
                    try:
                        step()
                    except Exception:
                        pass

    def export_to_csv(self, result_id):
        if not self.snapshot or self.snapshot['id'] != result_id:
            return {'success': False, 'error': 'Result is no longer available. Run the query again.'}
        return {'success': True, 'send_file': self.send_file(self.snapshot['path'], remove=False), 'rows': self.snapshot['rows']}

    def export_to_tsv(self, result_id):
        snapshot = self.snapshot
        if not snapshot or snapshot['id'] != result_id:
            return {'success': False, 'error': 'Result is no longer available. Run the query again.'}
        too_large = {'success': False, 'error': f'This result is more than {self.CLIPBOARD_CHARS:,} characters, too large to paste. Use Download CSV instead.'}
        csv.field_size_limit(max(csv.field_size_limit(), self.CLIPBOARD_CHARS))
        output = io.StringIO()
        writer = csv.writer(output, delimiter='\t', lineterminator='\r\n')
        with open(snapshot['path'], encoding='utf-8', newline='') as source:
            try:
                for row in csv.reader(source):
                    writer.writerow(row)
                    if output.tell() > self.CLIPBOARD_CHARS:
                        return too_large
            except csv.Error:
                return too_large
        return {'success': True, 'content': output.getvalue(), 'rows': snapshot['rows']}

    def export_stages(self, stages, wanted, write):
        self.check_stages(stages)
        created, last = ([], len(stages) - 1)
        self.conn.execute('BEGIN TRANSACTION')
        try:
            for index, (name, sql) in enumerate(stages):
                self.require_single_select(sql, 'A recipe runs as one SELECT query per step. Use the SQL workbench for other statements.')
                if name is None or index == last:
                    source = sql
                else:
                    quoted = '"' + name + '"'
                    self.conn.execute(f'CREATE OR REPLACE TEMP TABLE {quoted} AS {sql}')
                    created.append(quoted)
                    source = f'SELECT * FROM temp.main.{quoted}'
                if index in wanted:
                    write(index, source)
                while len(created) > 1:
                    self.conn.execute(f'DROP TABLE IF EXISTS temp.main.{created.pop(0)}')
        finally:
            try:
                self.conn.execute('ROLLBACK')
            except Exception:
                pass

    def export_recipe_workbook(self, stages, sheets):
        if not isinstance(sheets, list) or not sheets or (not isinstance(stages, list)) or (not all((isinstance(item, list) and len(item) == 2 and (type(item[0]) is int) and (0 <= item[0] < len(stages)) and isinstance(item[1], str) and (0 < len(item[1]) <= 31) and (not re.search('[\\[\\]:*?/\\\\]', item[1])) for item in sheets))) or ([item[0] for item in sheets] != sorted({item[0] for item in sheets})) or (len({item[1].casefold() for item in sheets}) != len(sheets)):
            return {'success': False, 'error': 'Choose at least one step, each with its own sheet name.'}
        names, rows, book = (dict(sheets), [], None)
        try:
            from openpyxl import Workbook
            from openpyxl.cell import WriteOnlyCell
            from openpyxl.cell.cell import ERROR_CODES, ILLEGAL_CHARACTERS_RE
            book = Workbook(write_only=True)

            def write(index, sql):
                name = names[index]
                cursor = self.conn.execute(sql)
                columns = [item[0] for item in cursor.description]
                if len(columns) > self.XLSX_COLUMNS:
                    raise ValueError(f'Sheet “{name}” has {len(columns):,} columns; one Excel sheet holds {self.XLSX_COLUMNS:,}. Leave that step out or use Download CSV.')
                sheet = book.create_sheet(name)
                cell = lambda value: self.excel_cell(sheet, value, WriteOnlyCell, ILLEGAL_CHARACTERS_RE, ERROR_CODES)
                sheet.append([cell(str(column)) for column in columns])
                count = 0
                while True:
                    batch = cursor.fetchmany(512)
                    if not batch:
                        break
                    count += len(batch)
                    if count > self.XLSX_ROWS:
                        raise ValueError(f'Sheet “{name}” has more than {self.XLSX_ROWS:,} rows, more than one Excel sheet holds. Leave that step out or use Download CSV.')
                    for row in batch:
                        sheet.append([cell(value) for value in row])
                rows.append(count)
            self.export_stages(stages, names, write)
            book.save(self.WORKBOOK_PATH)
            return {'success': True, 'send_file': self.send_file(self.WORKBOOK_PATH, remove=True), 'rows': rows}
        except Exception as error:
            self.discard_workbook(book)
            return {'success': False, 'error': str(error) or f'{type(error).__name__} while building the workbook. Use Download CSV instead.'}

    def export_recipe_parquet(self, stages):
        path = self.PARQUET_PATH
        try:
            described = {}

            def write(index, sql):
                wanted = self.conn.execute(f'DESCRIBE {sql}').fetchall()
                self.conn.execute(f"COPY ({sql}) TO '{path}' (FORMAT parquet)")
                columns = self.conn.execute(f"DESCRIBE SELECT * FROM read_parquet('{path}')").fetchall()
                described.update(columns=[item[0] for item in columns], types=[item[1] for item in columns], rows=self.conn.execute(f"SELECT count(*) FROM read_parquet('{path}')").fetchone()[0], changed=[[before[0], before[1], after[1]] for before, after in zip(wanted, columns) if before[1] != after[1]])
            if not isinstance(stages, list) or not stages:
                raise ValueError('A staged recipe needs one name and query per step.')
            self.export_stages(stages, {len(stages) - 1}, write)
            return {'success': True, 'send_file': self.send_file(path, remove=True), **described}
        except Exception as error:
            try:
                if os.path.exists(path):
                    os.remove(path)
            except OSError:
                pass
            return {'success': False, 'error': str(error)}
csv_tools = DuckDBCSVTools()

def worker_call(method, args_json):
    allowed = {'get_version', 'load_csv_file', 'load_excel_file', 'load_additional_sheet', 'reload_table_with_headers', 'replace_file', 'drop_table', 'get_table_info', 'preview_query', 'execute_query', 'execute_recipe', 'execute_recipe_stages', 'export_to_csv', 'export_to_tsv', 'check_export_to_xlsx', 'export_to_xlsx', 'export_recipe_workbook', 'export_recipe_parquet'}
    if method not in allowed:
        raise ValueError('Unknown engine operation')
    result = getattr(csv_tools, method)(*json.loads(args_json))
    return json.dumps(result, allow_nan=False)
